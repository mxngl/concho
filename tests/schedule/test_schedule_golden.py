"""Golden test: the Island 2026 schedule outputs of all 15 steps (P2.6, P3B.8).

Re-runs every migrated step (``python -m engines.schedule <step>``) on the IPD_Challenge@989a6b7
inputs and compares each output file with ``tests/fixtures/schedule_golden.json``:

- ``sha256`` after the masking of the P1.7 equivalence test (run-root paths, random P6 GUIDs,
  the relative FBX link); xlsx files are hashed by cell values, PNG charts only have to exist;
- readable ``metrics`` (row counts, first/last dates, durations, task/zone counts, delivery
  counts per window), so a failure shows *what* changed before the hash does.

Unlike the equivalence test, this does not need the original scripts: the reference is the
JSON file, so it keeps working once IPD_Challenge is archived (only its inputs are read). The
team rules and BIM map come from ``engines/schedule/examples/island/`` (byte-identical to the
989a6b7 copies), so edits there are caught too. The arguments are those of the equivalence
test (``_migrated_args``), incl. the committed Manufacton workbooks for ``delivery-windows``,
because ``manufacton-orders`` fails on this data (its error line is pinned instead).

The regenerated ``Micro_Schedule.csv`` (6,625 rows, 2029-10-01 09:00 -> 2030-03-15 19:40) is
the golden micro schedule. The committed one in IPD_Challenge is stale; that stays a strict
xfail in ``test_schedule_equivalence.py``.

Skipped without the IPD_Challenge fixture (fails in the CI job ``reference``). After an
*intended* output change, regenerate the JSON and review its diff::

    python tests/schedule/test_schedule_golden.py --update
"""

from __future__ import annotations

import hashlib
import json
import re
import sys
import tempfile
from pathlib import Path

import pytest
from test_schedule_equivalence import _FBX_PATH, _GUID, _migrated_args, _run

REPO_ROOT = Path(__file__).resolve().parents[2]
GOLDEN_PATH = REPO_ROOT / "tests" / "fixtures" / "schedule_golden.json"
EXAMPLES = REPO_ROOT / "engines" / "schedule" / "examples" / "island"
DATE_FMT = "%Y-%m-%d %H:%M"
SUPERSTRUCTURE = {"A1330", "A1340", "A1350", "A1420"}  # columns, floor, beams, roof
BUFFER_TASK = "A1880"  # Hurricane Contingency Buffer


def golden_args(ipd: Path, out: Path) -> dict[str, list[str]]:
    """Equivalence-test arguments, with the committed Island rules and BIM map."""
    steps = _migrated_args(ipd, out)
    micro = steps["micro-schedule"]
    micro[micro.index("--bim-map") + 1] = str(EXAMPLES / "ALICE_BIM_Map.csv")
    micro[micro.index("--rules") + 1] = str(EXAMPLES / "micro_schedule_rules.json")
    return steps


def run_pipeline(ipd: Path, out: Path) -> dict[str, dict[str, object]]:
    """Run all 15 steps into ``out``; return exit code and last stderr line per step."""
    steps = {}
    for step, args in golden_args(ipd, out).items():
        proc = _run(["-m", "engines.schedule", step, *args], cwd=out.parent)
        entry: dict[str, object] = {"returncode": proc.returncode}
        if proc.returncode != 0:
            lines = proc.stderr.strip().splitlines()
            entry["error"] = _mask(lines[-1] if lines else "", [out, ipd])
        steps[step] = entry
    return steps


# --- masking and hashing -------------------------------------------------------------------


def _mask(text: str, roots: list[Path]) -> str:
    for root in roots:
        text = text.replace(str(root), "<ROOT>")
    text = _GUID.sub("{GUID}", text)
    return _FBX_PATH.sub('"<FBX>"', text)


def _xlsx_values(path: Path) -> list[list[list[object]]]:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        return [[list(row) for row in sheet.iter_rows(values_only=True)] for sheet in workbook]
    finally:
        workbook.close()


def masked_sha256(path: Path, roots: list[Path]) -> str | None:
    if path.suffix == ".png":
        return None  # matplotlib rendering; only existence is checked
    if path.suffix == ".xlsx":
        data = json.dumps(_xlsx_values(path), default=str).encode("utf-8")
    else:
        data = _mask(path.read_text(encoding="utf-8"), roots).encode("utf-8")
    return hashlib.sha256(data).hexdigest()


# --- readable metrics ----------------------------------------------------------------------


def _round(value: float) -> float:
    return round(float(value), 4)


def _ts(series):
    import pandas as pd

    return pd.to_datetime(series, format="ISO8601")


def _span(start, end) -> dict[str, object]:
    import pandas as pd

    return {
        "first_start": start.strftime(DATE_FMT),
        "last_end": end.strftime(DATE_FMT),
        "span_calendar_days": (end.normalize() - start.normalize()).days + 1,
        "span_weekdays": len(pd.bdate_range(start.normalize(), end.normalize())),
    }


def _macro(df) -> dict[str, object]:
    start, end = _ts(df["start_date"]), _ts(df["end_date"])
    return {"tasks": int(df["task_id"].nunique()), **_span(start.min(), end.max())}


def _micro(df) -> dict[str, object]:
    start, end = _ts(df["element_start"]), _ts(df["element_end"])
    superstructure = df["task_id"].isin(SUPERSTRUCTURE)
    no_buffer = df["task_id"] != BUFFER_TASK
    windows = (
        df.assign(start=start, end=end)
        .groupby("task_id")
        .agg(start=("start", "min"), end=("end", "max"))
    )
    return {
        "rows": len(df),
        "tasks": int(df["task_id"].nunique()),
        "micro_task_ids": int(df["micro_task_id"].nunique()),
        "bim_elements": int(df.loc[df["source_model"] != "Non-BIM", "element_key"].nunique()),
        **_span(start.min(), end.max()),
        "end_without_buffer": end[no_buffer].max().strftime(DATE_FMT),
        "superstructure_start": start[superstructure].min().strftime(DATE_FMT),
        "superstructure_end": end[superstructure].max().strftime(DATE_FMT),
        "final_inspections_end": end[df["task_id"] == "A1860"].max().strftime(DATE_FMT),
        "scheduled_hours_total": _round(df["scheduled_duration_hr"].sum()),
        "task_windows": {
            task: [group_start.strftime(DATE_FMT), group_end.strftime(DATE_FMT)]
            for task, group_start, group_end in zip(
                windows.index, windows["start"], windows["end"], strict=True
            )
        },
    }


def _macro_view(df) -> dict[str, object]:
    return {
        "activities": len(df),
        "tasks": int(df["task_id"].nunique()),
        "first_start": _ts(df["start"]).min().strftime(DATE_FMT),
        "last_finish": _ts(df["finish"]).max().strftime(DATE_FMT),
        "bim_elements": int(df.loc[df["level"] != "Basement / Site", "element_count"].sum()),
    }


def _takt_schedule(df) -> dict[str, object]:
    return {
        "rows": len(df),
        "zones": int(df["takt_zone_id"].nunique()),
        "total_hours": _round(df["finish_hour"].max()),
        "first_start": _ts(df["start"]).min().strftime(DATE_FMT),
        "last_finish": _ts(df["finish"]).max().strftime(DATE_FMT),
    }


def _crew_utilization(df) -> dict[str, object]:
    return {
        "rows": len(df),
        "utilization": {row.crew: _round(row.utilization) for row in df.itertuples()},
    }


def _order_windows(df) -> dict[str, object]:
    windows = {}
    for window, group in df.groupby("window", sort=False):
        counts = group["production_order_count"]
        windows[window] = {
            "delivery_days": int((counts > 0).sum()),
            "peak_orders_per_delivery": int(counts.max()),
            "orders": int(counts.sum()),
        }
    return {"rows": len(df), "windows": windows}


def _delivery_summary(df) -> dict[str, object]:
    cols = ["total_delivered_elements", "total_delivered_cf", "max_single_delivery_cf",
            "max_peak_on_site_cf", "max_peak_on_site_elements"]
    return {
        "rows": len(df),
        "windows": {row["window"]: {c: _round(row[c]) for c in cols} for _, row in df.iterrows()},
    }


def _delivery_units(df) -> dict[str, object]:
    return {
        "rows": len(df),
        "by_source": {k: int(v) for k, v in df["source"].value_counts().sort_index().items()},
        "elements": int(df["element_count"].sum()),
        "first_need_date": str(df["need_date"].min()),
        "last_need_date": str(df["need_date"].max()),
    }


CSV_METRICS = {
    "Macro_Schedule.csv": _macro,
    "Micro_Schedule.csv": _micro,
    "ALICE_Task_Schedule_Macro_View.csv": _macro_view,
    "Takt_Schedule.csv": _takt_schedule,
    "Takt_Zones.csv": lambda df: {"zones": len(df), "rooms": int(df["room_count"].sum())},
    "Takt_Crew_Idle_Report.csv": _crew_utilization,
    "production_order_count_by_delivery_window.csv": _order_windows,
    "delivery_window_summary_metrics.csv": _delivery_summary,
    "delivery_units_by_micro_schedule.csv": _delivery_units,
    "central_bim_model_with_takt.csv": lambda df: {
        "rows": len(df), "with_takt_id": int(df["takt_id"].notna().sum())},
    "Prefab_Wall_Mapping.csv": lambda df: {
        "rows": len(df), "prefab_groups": int(df["prefab_group_id"].nunique())},
    "Parts_Summary.csv": lambda df: {
        "parts": len(df), "elements": int(df["element_count"].sum())},
}


def metrics(path: Path) -> dict[str, object]:
    if path.suffix == ".csv":
        import pandas as pd

        df = pd.read_csv(path, low_memory=False)
        base = {"rows": len(df), "columns": len(df.columns)}
        return {**base, **CSV_METRICS.get(path.name, lambda _: {})(df)}
    if path.suffix == ".xlsx":
        return {"sheet_rows": [len(sheet) for sheet in _xlsx_values(path)]}
    if path.suffix == ".xml":
        text = path.read_text(encoding="utf-8")
        return {"activities": len(re.findall(r"<Activity>", text))}
    if path.suffix in (".md", ".html"):
        return {"lines": len(path.read_text(encoding="utf-8").splitlines())}
    return {}


def collect(out: Path, ipd: Path, steps: dict[str, dict[str, object]]) -> dict[str, object]:
    """Golden record of a pipeline run in ``out``."""
    files = {}
    for path in sorted(p for p in out.rglob("*") if p.is_file()):
        entry: dict[str, object] = {}
        sha = masked_sha256(path, [out, ipd])
        if sha is not None:
            entry["sha256"] = sha
        entry["metrics"] = metrics(path)
        files[path.relative_to(out).as_posix()] = entry
    return {"steps": steps, "files": files}


def generate(ipd: Path, workdir: Path) -> dict[str, object]:
    out = workdir / "schedule_golden"
    return collect(out, ipd, run_pipeline(ipd, out))


# --- tests ---------------------------------------------------------------------------------


# Missing only while bootstrapping with ``--update``.
GOLDEN = (
    json.loads(GOLDEN_PATH.read_text(encoding="utf-8"))
    if GOLDEN_PATH.exists() else {"steps": {}, "files": {}}
)


@pytest.fixture(scope="module")
def actual(tmp_path_factory: pytest.TempPathFactory, ipd_challenge_dir: Path) -> dict:
    return generate(ipd_challenge_dir, tmp_path_factory.mktemp("schedule_golden"))


@pytest.mark.parametrize("step", list(GOLDEN["steps"]))
def test_step_result_matches_golden(actual: dict, step: str) -> None:
    assert actual["steps"][step] == GOLDEN["steps"][step]


def test_output_files_match_golden(actual: dict) -> None:
    assert sorted(actual["files"]) == sorted(GOLDEN["files"])


def _diff(expected: object, got: object, key: str = "") -> list[str]:
    """Flat list of ``key: golden -> actual`` lines for nested metric dicts."""
    if isinstance(expected, dict) and isinstance(got, dict):
        lines = []
        for sub in sorted(set(expected) | set(got)):
            lines += _diff(expected.get(sub), got.get(sub), f"{key}.{sub}" if key else sub)
        return lines
    return [] if expected == got else [f"{key}: {expected!r} -> {got!r}"]


@pytest.mark.parametrize("name", list(GOLDEN["files"]))
def test_output_metrics_match_golden(actual: dict, name: str) -> None:
    assert name in actual["files"], f"{name} was not written"
    diff = _diff(GOLDEN["files"][name]["metrics"], actual["files"][name]["metrics"])
    assert not diff, "metrics differ (golden -> actual):\n" + "\n".join(diff)


@pytest.mark.parametrize(
    "name", [name for name, entry in GOLDEN["files"].items() if "sha256" in entry]
)
def test_output_sha256_matches_golden(actual: dict, name: str) -> None:
    assert name in actual["files"], f"{name} was not written"
    assert actual["files"][name].get("sha256") == GOLDEN["files"][name]["sha256"]


def test_current_micro_schedule_is_golden_reference(actual: dict) -> None:
    """The regenerated micro schedule replaces the stale committed one as reference."""
    micro = GOLDEN["files"]["micro/Micro_Schedule.csv"]["metrics"]
    assert micro["rows"] == 6625
    assert (micro["first_start"], micro["last_end"]) == ("2029-10-01 09:00", "2030-03-15 19:40")
    name = "micro/Micro_Schedule.csv"
    assert actual["files"][name] == GOLDEN["files"][name]


def test_takt_plan_level_1_reference(actual: dict) -> None:
    """Roadmap P2.6: takt planner Level 1 has 16 zones and 192.36 working hours."""
    takt = actual["files"]["takt/Takt_Schedule.csv"]["metrics"]
    assert (takt["zones"], takt["total_hours"]) == (16, 192.36)


def main(argv: list[str]) -> int:
    """``--update``: regenerate ``schedule_golden.json`` from the fixture inputs."""
    if argv != ["--update"]:
        print(__doc__)
        return 2
    sys.path.insert(0, str(REPO_ROOT / "tests"))
    from conftest import reference_repo

    ipd = reference_repo("IPD_Challenge", override_env="IPD_CHALLENGE_DIR")
    if ipd is None:
        print("IPD_Challenge fixture not found (python scripts/fetch_fixtures.py)")
        return 1
    with tempfile.TemporaryDirectory() as tmp:
        golden = generate(ipd, Path(tmp))
    golden = {
        "_comment": [
            "P2.6 golden record of the Island 2026 schedule outputs (all 15 steps).",
            "Inputs: IPD_Challenge@989a6b7 + engines/schedule/examples/island. No output files are",
            "committed; sha256 is taken after masking run-root paths, P6 GUIDs and the FBX link",
            "(xlsx: cell values; png: existence only). Regenerate with",
            "python tests/schedule/test_schedule_golden.py --update, then review the diff.",
        ],
        **golden,
    }
    text = json.dumps(golden, indent=2, ensure_ascii=False) + "\n"
    GOLDEN_PATH.write_text(text, encoding="utf-8")
    print(f"wrote {GOLDEN_PATH.relative_to(REPO_ROOT)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv[1:]))
