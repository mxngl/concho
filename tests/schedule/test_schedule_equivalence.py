"""Equivalence test: migrated schedule CLIs vs. the original IPD_Challenge scripts.

Uses the IPD_Challenge@989a6b7 checkout from ``IPD_CHALLENGE_DIR`` if set, else
``IPD_Challenge`` in the shared fixture root (``CONCHO_FIXTURES_DIR`` or ``.fixtures/``, see
``tests/conftest.py``); skipped without it (fails in the CI job ``reference``). Nothing from
that checkout is copied into this repo.

The original scripts run, in pipeline order, inside a temporary copy of the checkout (they
write next to themselves). The migrated steps run via ``python -m engines.schedule <step>``
on the same inputs: raw inputs are read from ``IPD_CHALLENGE_DIR``, intermediate files from
the *original* run (P3B.8: each step is compared in isolation, so an intended fix only
changes the outputs of the step it fixes; the chained migrated run is covered by
``test_schedule_golden.py``). Where the original read a file that is committed in the checkout
but not produced earlier in the chain (``takt_zones.json``, ``room_takt_zones.csv``, the
takt productivity rates, the previous ``Revit_Assembly_Id_Map.csv``, the Manufacton order
workbooks), the migrated step gets the committed file too.

Compared per output: text files byte-for-byte after replacing both run roots with a
placeholder; P6 XML with GUIDs masked; xlsx by cell values; PNGs are not compared
(rendering). ``manufacton-orders`` fails in both implementations on this data with the same
error (see engines/schedule/README.md), so only the error is compared.

Also checks the reference sha256 values of docs/ROADMAP.md §1 (P0.4).
"""

from __future__ import annotations

import csv
import hashlib
import io
import os
import re
import shutil
import subprocess
import sys
from pathlib import Path

import pytest

PE = "src/Planning_engine"
ALICE_OUT = f"{PE}/ALICE_BIM_mapper/outputs"
MICRO_OUT = f"{PE}/Micro_Schedule_Generator/outputs"
PREFAB_OUT = f"{PE}/Prefab_BIM_Mapper/outputs"
FUZOR_OUT = f"{PE}/Fuzor_Mapper/outputs"
TAKT_OUT = "src/Takt_engine/outputs"
ZONES_OUT = "outputs/takt_zones"
DELIVERY_OUT = "outputs/delivery_window_analysis"
ROOMS_OUT = "outputs/room_boundaries"

# sha256 of the committed Island reference outputs (docs/ROADMAP.md §1, P0.4).
REFERENCE_SHA256 = {
    f"{ALICE_OUT}/Macro_Schedule.csv":
        "d267a913b33513586eba4a8a8b7128fce9122b816f51a6435399631ffad529b2",
    f"{MICRO_OUT}/Micro_Schedule.csv":
        "7841a4b740587d41149227e92b3f6f76f276c2556350e102735ad0135511b7f4",
    f"{TAKT_OUT}/Takt_Schedule.csv":
        "17fa8009033019405aff0fda921f011bb0f821426a2c4d5eb9042e0a60b8379b",
    f"{ZONES_OUT}/central_bim_model_with_takt.csv":
        "ff008794690305678f0eaf09b9dc79e053fe40751713f2f90abacc8118d7b09b",
}

# The Island takt plan in IPD_Challenge was generated with two rooms per takt zone.
TAKT_ARGS = ["--level", "L 1", "--rooms-per-zone", "2"]

# Original script (relative to the checkout) per migrated step, in pipeline order.
ORIGINAL_SCRIPTS = {
    "takt-zones": ["src/takt_zone_calibrator.py"],
    "llm-context": [f"{PE}/generate_llm_bim_context.py"],
    "alice-inputs": [f"{PE}/ALICE_BIM_mapper/generate_inputs.py"],
    "prefab-walls": [f"{PE}/Prefab_BIM_Mapper/generate_prefab_wall_mapping.py"],
    "micro-schedule": [f"{PE}/Micro_Schedule_Generator/generate_micro_schedule.py"],
    "alice-p6-xml": [f"{PE}/ALICE_BIM_mapper/generate_p6_task_schedule_xml.py"],
    "fuzor-xml": [f"{PE}/Fuzor_Mapper/generate_fuzor_p6_xml.py"],
    "manufacton-parts": [f"{PE}/Prefab_BIM_Mapper/generate_parts_import.py"],
    "manufacton-assemblies": [f"{PE}/Prefab_BIM_Mapper/generate_assembly_import.py"],
    "manufacton-orders": [f"{PE}/Prefab_BIM_Mapper/generate_kit_import.py"],
    "delivery-windows": [f"{PE}/Logistics_Analysis/compare_delivery_windows.py"],
    "takt-plan": ["-m", "src.Takt_engine.takt_planner", *TAKT_ARGS],
    "takt-viewer": [f"{PE}/Micro_Schedule_Generator/generate_takt_viewer.py"],
    "spatial-viewer": [f"{PE}/generate_spatial_visualizer.py"],
}

# (step, original output relative to the checkout, migrated output file name)
COMPARED_OUTPUTS = [
    ("takt-zones", f"{ZONES_OUT}/central_bim_model.csv", "central_bim_model.csv"),
    ("takt-zones", f"{ZONES_OUT}/central_bim_model_with_takt.csv",
     "central_bim_model_with_takt.csv"),
    ("llm-context", f"{ZONES_OUT}/central_bim_model_llm_context.csv",
     "central_bim_model_llm_context.csv"),
    ("alice-inputs", f"{ALICE_OUT}/Macro_Schedule.csv", "Macro_Schedule.csv"),
    ("alice-inputs", f"{ALICE_OUT}/Crew.csv", "Crew.csv"),
    ("alice-inputs", f"{ALICE_OUT}/Equipment.csv", "Equipment.csv"),
    ("alice-inputs", f"{ALICE_OUT}/Tasks.csv", "Tasks.csv"),
    ("alice-inputs", f"{ALICE_OUT}/Missing_data.md", "Missing_data.md"),
    ("prefab-walls", f"{PREFAB_OUT}/Prefab_Wall_Mapping.csv", "Prefab_Wall_Mapping.csv"),
    ("micro-schedule", f"{MICRO_OUT}/Micro_Schedule.csv", "Micro_Schedule.csv"),
    ("micro-schedule", f"{MICRO_OUT}/Micro_Schedule_Log.md", "Micro_Schedule_Log.md"),
    ("alice-p6-xml", f"{ALICE_OUT}/ALICE_Task_Schedule.xml", "ALICE_Task_Schedule.xml"),
    ("alice-p6-xml", f"{ALICE_OUT}/ALICE_Task_Schedule_Macro_View.csv",
     "ALICE_Task_Schedule_Macro_View.csv"),
    ("fuzor-xml", f"{FUZOR_OUT}/Fuzor_Micro_Schedule.xml", "Fuzor_Micro_Schedule.xml"),
    ("fuzor-xml", f"{FUZOR_OUT}/Revit_4D_Build_Code_Map.csv", "Revit_4D_Build_Code_Map.csv"),
    ("manufacton-parts", f"{PREFAB_OUT}/Parts_Import.xlsx", "Parts_Import.xlsx"),
    ("manufacton-parts", f"{PREFAB_OUT}/Parts_Import.csv", "Parts_Import.csv"),
    ("manufacton-parts", f"{PREFAB_OUT}/Parts_Summary.csv", "Parts_Summary.csv"),
    ("manufacton-assemblies", f"{PREFAB_OUT}/Assembly_Import.xlsx", "Assembly_Import.xlsx"),
    ("delivery-windows", f"{DELIVERY_OUT}/delivery_units_by_micro_schedule.csv",
     "delivery_units_by_micro_schedule.csv"),
    ("delivery-windows", f"{DELIVERY_OUT}/delivery_window_daily_timeseries.csv",
     "delivery_window_daily_timeseries.csv"),
    ("delivery-windows", f"{DELIVERY_OUT}/delivery_window_summary_metrics.csv",
     "delivery_window_summary_metrics.csv"),
    ("delivery-windows", f"{DELIVERY_OUT}/production_order_count_by_delivery_window.csv",
     "production_order_count_by_delivery_window.csv"),
    ("takt-plan", f"{TAKT_OUT}/Takt_Zones.csv", "Takt_Zones.csv"),
    ("takt-plan", f"{TAKT_OUT}/Takt_Element_Allocations.csv", "Takt_Element_Allocations.csv"),
    ("takt-plan", f"{TAKT_OUT}/Takt_Element_Splits.csv", "Takt_Element_Splits.csv"),
    ("takt-plan", f"{TAKT_OUT}/Takt_Schedule.csv", "Takt_Schedule.csv"),
    ("takt-plan", f"{TAKT_OUT}/Takt_Crew_Idle_Report.csv", "Takt_Crew_Idle_Report.csv"),
    ("takt-plan", f"{TAKT_OUT}/Takt_Equipment_Inputs.csv", "Takt_Equipment_Inputs.csv"),
    ("takt-plan", f"{TAKT_OUT}/Takt_Productivity_Rates.csv", "Takt_Productivity_Rates.csv"),
    ("takt-plan", f"{TAKT_OUT}/Takt_Report.md", "Takt_Report.md"),
    ("takt-plan", f"{TAKT_OUT}/Takt_Planner.html", "Takt_Planner.html"),
    ("takt-plan", f"{TAKT_OUT}/Takt_Model_Viewer.html", "Takt_Model_Viewer.html"),
    ("takt-viewer", f"{MICRO_OUT}/Micro_Schedule_Takt_Viewer.html",
     "Micro_Schedule_Takt_Viewer.html"),
    ("spatial-viewer", f"{PE}/spatial_visualizer_micro.html", "spatial_visualizer_micro.html"),
]

# Outputs that a P3B.8 bug fix changes on purpose (see engines/schedule/README.md, "Fixed in
# P3B.8"). They are checked by a dedicated function instead of byte equality.
P3B8_CHANGED_OUTPUTS = {
    ("takt-zones", "central_bim_model_with_takt.csv"): "_check_takt_ids_fix",
}

# P3B.8 fix 1 (takt zones keep their last corner): changed takt_id values in the Island model.
TAKT_FIX_NEWLY_ASSIGNED = 900  # no takt zone before, one now
TAKT_FIX_MOVED = 1  # other zone (element 1293125, L 1 Zone 5 -> L 1 Zone 1)

_GUID = re.compile(r"\{[0-9A-F]{8}-[0-9A-F]{4}-[0-9A-F]{4}-[0-9A-F]{4}-[0-9A-F]{12}\}")
# Takt_Model_Viewer.html links the FBX relative to its own folder, which differs per layout.
_FBX_PATH = re.compile(r'"[^"\n]*\.fbx"', re.IGNORECASE)


def _env() -> dict[str, str]:
    env = dict(os.environ)
    env["MPLBACKEND"] = "Agg"
    env["PYTHONIOENCODING"] = "utf-8"
    return env


def _run(args: list[str], cwd: Path) -> subprocess.CompletedProcess[str]:
    return subprocess.run(
        [sys.executable, *args], cwd=cwd, env=_env(), capture_output=True, text=True,
        encoding="utf-8", errors="replace", check=False,
    )


def _migrated_args(ipd: Path, new: Path, inputs_from: Path | None = None) -> dict[str, list[str]]:
    """CLI arguments of every migrated step, fed like the original run in the checkout.

    By default each step reads the outputs of the previous *migrated* step (a chained run,
    used by the golden test). With ``inputs_from`` (the original run's checkout copy), every
    step reads the intermediate files of the *original* run instead, so each step is
    compared in isolation and an intended change (P3B.8) only shows in its own outputs.
    Outputs always go to ``new``.
    """
    alice, micro, prefab, fuzor = new / "alice", new / "micro", new / "prefab", new / "fuzor"
    zones, rooms = new / "takt_zones", new / "rooms"
    if inputs_from is not None:
        alice, micro, prefab, fuzor = (
            inputs_from / ALICE_OUT, inputs_from / MICRO_OUT, inputs_from / PREFAB_OUT,
            inputs_from / FUZOR_OUT,
        )
        zones, rooms = inputs_from / ZONES_OUT, inputs_from / ROOMS_OUT
    out = {
        "zones": new / "takt_zones", "alice": new / "alice", "prefab": new / "prefab",
        "micro": new / "micro", "fuzor": new / "fuzor", "rooms": new / "rooms",
    }
    committed_prefab = ipd / PREFAB_OUT
    return {
        "takt-zones": [
            "--schedules-dir", str(ipd / "revit_schedules" / "Current"),
            "--takt-zones", str(ipd / ZONES_OUT / "takt_zones.json"),
            "--out-dir", str(out["zones"]),
        ],
        "llm-context": [
            "--central-bim-with-takt", str(zones / "central_bim_model_with_takt.csv"),
            "--out-dir", str(out["zones"]),
        ],
        "alice-inputs": [
            "--workbook", str(ipd / PE / "ALICE_BIM_mapper" / "inputs" / "ALICE_macro.xlsx"),
            "--out-dir", str(out["alice"]),
        ],
        "prefab-walls": [
            "--central-bim-with-takt", str(zones / "central_bim_model_with_takt.csv"),
            "--out-dir", str(out["prefab"]),
        ],
        "micro-schedule": [
            "--macro-schedule", str(alice / "Macro_Schedule.csv"),
            "--tasks", str(alice / "Tasks.csv"),
            "--crew", str(alice / "Crew.csv"),
            "--equipment", str(alice / "Equipment.csv"),
            "--bim-map",
            str(ipd / PE / "Micro_Schedule_Generator" / "inputs" / "ALICE_BIM_Map.csv"),
            "--central-bim-with-takt", str(zones / "central_bim_model_with_takt.csv"),
            "--llm-context", str(zones / "central_bim_model_llm_context.csv"),
            "--prefab-wall-mapping", str(prefab / "Prefab_Wall_Mapping.csv"),
            "--rules", str(
                ipd / PE / "Micro_Schedule_Generator" / "inputs" / "micro_schedule_rules.json"
            ),
            "--room-boundaries", str(_room_boundaries(ipd)),
            "--out-dir", str(out["micro"]),
        ],
        "alice-p6-xml": [
            "--micro-schedule", str(micro / "Micro_Schedule.csv"),
            "--workbook", str(ipd / PE / "ALICE_BIM_mapper" / "inputs" / "ALICE_macro.xlsx"),
            "--out-dir", str(out["alice"]),
        ],
        "fuzor-xml": [
            "--micro-schedule", str(micro / "Micro_Schedule.csv"),
            "--tasks", str(alice / "Tasks.csv"),
            "--crew", str(alice / "Crew.csv"),
            "--equipment", str(alice / "Equipment.csv"),
            "--out-dir", str(out["fuzor"]),
        ],
        "manufacton-parts": [
            "--template", str(ipd / PE / "Prefab_BIM_Mapper" / "inputs" / "Parts Import.xlsx"),
            "--central-bim-with-takt", str(zones / "central_bim_model_with_takt.csv"),
            "--micro-schedule", str(micro / "Micro_Schedule.csv"),
            "--assembly-id-map", str(committed_prefab / "Revit_Assembly_Id_Map.csv"),
            "--out-dir", str(out["prefab"]),
        ],
        "manufacton-assemblies": [
            "--template",
            str(ipd / PE / "Prefab_BIM_Mapper" / "inputs" / "Assembly.Import.xlsx"),
            "--parts-import", str(prefab / "Parts_Import.xlsx"),
            "--parts-summary", str(prefab / "Parts_Summary.csv"),
            "--micro-schedule", str(micro / "Micro_Schedule.csv"),
            "--central-bim-with-takt", str(zones / "central_bim_model_with_takt.csv"),
            "--build-code-map", str(fuzor / "Revit_4D_Build_Code_Map.csv"),
            "--out-dir", str(out["prefab"]),
        ],
        "manufacton-orders": [
            "--order-template",
            str(ipd / PE / "Prefab_BIM_Mapper" / "inputs" / "ORDER IMPORT TEMPLATE.xlsx"),
            "--item-template",
            str(ipd / PE / "Prefab_BIM_Mapper" / "inputs" / "ITEM IMPORT TEMPLATE.xlsx"),
            "--vendors", str(ipd / PE / "Prefab_BIM_Mapper" / "inputs" / "vendors.csv"),
            "--mapping", str(
                ipd / PE / "Prefab_BIM_Mapper" / "inputs"
                / "4d_build_code_to_assembly_id_mapping.csv"
            ),
            "--assembly-import", str(prefab / "Assembly_Import.xlsx"),
            "--parts-summary", str(prefab / "Parts_Summary.csv"),
            "--build-code-map", str(fuzor / "Revit_4D_Build_Code_Map.csv"),
            "--micro-schedule", str(micro / "Micro_Schedule.csv"),
            "--llm-context", str(zones / "central_bim_model_llm_context.csv"),
            "--out-dir", str(out["prefab"]),
        ],
        # manufacton-orders fails on this data (see module docstring), so the original
        # delivery analysis read the committed Manufacton workbooks and maps.
        "delivery-windows": [
            "--micro-schedule", str(micro / "Micro_Schedule.csv"),
            "--llm-context", str(zones / "central_bim_model_llm_context.csv"),
            "--production-order", str(committed_prefab / "Production_Order.xlsx"),
            "--production-order-items", str(committed_prefab / "Production_Order_Items.xlsx"),
            "--kit-map", str(committed_prefab / "Revit_Kit_Parameter_Map.csv"),
            "--assembly-map", str(committed_prefab / "Revit_Assembly_Id_Map.csv"),
            "--out-dir", str(new / "delivery"),
        ],
        # P3B.8 fix 2: generator for room_takt_zones.csv (no original script). In the isolated
        # comparison the takt plan reads the committed file, as the original did.
        "room-takt-zones": [
            "--room-boundaries", str(_room_boundaries(ipd)),
            "--out-dir", str(out["rooms"]),
        ],
        "takt-plan": [
            *TAKT_ARGS,
            "--central-bim", str(zones / "central_bim_model.csv"),
            "--room-takt-zones", str(rooms / "room_takt_zones.csv"),
            "--crew", str(alice / "Crew.csv"),
            "--equipment", str(alice / "Equipment.csv"),
            "--productivity-rates", str(ipd / TAKT_OUT / "Takt_Productivity_Rates.csv"),
            "--room-boundaries", str(_room_boundaries(ipd)),
            "--fbx", str(_fbx(ipd)),
            "--out-dir", str(new / "takt"),
        ],
        "takt-viewer": [
            "--micro-schedule", str(micro / "Micro_Schedule.csv"),
            "--alice-workbook",
            str(ipd / PE / "ALICE_BIM_mapper" / "inputs" / "ALICE_macro.xlsx"),
            "--out-dir", str(out["micro"]),
        ],
        "spatial-viewer": [
            "--micro-schedule", str(micro / "Micro_Schedule.csv"),
            *_spatial_floor_plans(ipd),
            "--out-dir", str(new / "viewers"),
        ],
    }


def _room_boundaries(ipd: Path) -> Path:
    (path,) = (ipd / "revit_schedules").glob("*_Room_Boundaries.csv")
    return path


def _fbx(ipd: Path) -> Path:
    (path,) = (ipd / "revit_schedules").glob("*.fbx")
    return path


def _spatial_floor_plans(ipd: Path) -> list[str]:
    cropped = ipd / "floor_plans" / "cropped_png"
    args = []
    for level, token in [("L -1", "Level -1_"), ("L 0", "Level 0_"), ("L 1", "Level 1_")]:
        for path in sorted(cropped.glob("*_page_0_cropped.png")):
            if token in path.name:
                args += ["--floor-plan", f"{level}={path}"]
    return args


# Migrated output folder per step (relative to the migrated run root).
MIGRATED_DIRS = {
    "takt-zones": "takt_zones", "llm-context": "takt_zones", "alice-inputs": "alice",
    "prefab-walls": "prefab", "micro-schedule": "micro", "alice-p6-xml": "alice",
    "fuzor-xml": "fuzor", "manufacton-parts": "prefab", "manufacton-assemblies": "prefab",
    "manufacton-orders": "prefab", "delivery-windows": "delivery",
    "room-takt-zones": "rooms", "takt-plan": "takt",
    "takt-viewer": "micro", "spatial-viewer": "viewers",
}


@pytest.fixture(scope="module")
def runs(
    tmp_path_factory: pytest.TempPathFactory, ipd_challenge_dir: Path
) -> dict[str, object]:
    """Run the original pipeline and the migrated pipeline once; return roots + results."""
    base = tmp_path_factory.mktemp("schedule_equivalence")
    orig = base / "original"
    new = base / "migrated"
    shutil.copytree(ipd_challenge_dir, orig, ignore=shutil.ignore_patterns(".git"))

    results: dict[str, dict[str, subprocess.CompletedProcess[str]]] = {"orig": {}, "new": {}}
    for step, script in ORIGINAL_SCRIPTS.items():
        results["orig"][step] = _run(script, cwd=orig)
    for step, args in _migrated_args(ipd_challenge_dir, new, inputs_from=orig).items():
        results["new"][step] = _run(["-m", "engines.schedule", step, *args], cwd=base)
    return {"orig": orig, "new": new, "ipd": ipd_challenge_dir, "results": results}


def _normalize(text: str, runs: dict[str, object]) -> str:
    for root in (runs["orig"], runs["ipd"], runs["new"]):
        text = text.replace(str(root), "<ROOT>")
    text = _GUID.sub("{GUID}", text)
    return _FBX_PATH.sub('"<FBX>"', text)


def _xlsx_values(path: Path) -> list[list[list[object]]]:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    return [[list(row) for row in sheet.iter_rows(values_only=True)] for sheet in workbook]


@pytest.mark.parametrize("step", list(ORIGINAL_SCRIPTS))
def test_step_exit_status_matches(runs: dict[str, object], step: str) -> None:
    orig = runs["results"]["orig"][step]
    new = runs["results"]["new"][step]
    if step == "manufacton-orders":
        # Known failure of the original on the 989a6b7 data; the migrated step must fail
        # identically.
        assert orig.returncode != 0 and new.returncode != 0
        assert orig.stderr.strip().splitlines()[-1] == new.stderr.strip().splitlines()[-1]
        return
    assert orig.returncode == 0, orig.stderr[-2000:]
    assert new.returncode == 0, new.stderr[-2000:]


@pytest.mark.parametrize(
    ("step", "original", "migrated"),
    COMPARED_OUTPUTS,
    ids=[f"{step}:{name}" for step, _, name in COMPARED_OUTPUTS],
)
def test_output_matches_original(
    runs: dict[str, object], step: str, original: str, migrated: str
) -> None:
    orig_path = runs["orig"] / original
    new_path = runs["new"] / MIGRATED_DIRS[step] / migrated
    assert new_path.exists(), f"migrated step {step} did not write {migrated}"
    if (step, migrated) in P3B8_CHANGED_OUTPUTS:
        check = globals()[P3B8_CHANGED_OUTPUTS[(step, migrated)]]
        check(*(_normalize(path.read_text(encoding="utf-8"), runs)
                for path in (orig_path, new_path)))
        return
    if migrated.endswith(".xlsx"):
        assert _xlsx_values(new_path) == _xlsx_values(orig_path)
        return
    orig_text = _normalize(orig_path.read_text(encoding="utf-8"), runs)
    new_text = _normalize(new_path.read_text(encoding="utf-8"), runs)
    assert new_text == orig_text


def _check_takt_ids_fix(orig_text: str, new_text: str) -> None:
    """P3B.8 fix 1: only ``takt_id`` differs, and only by the complete polygons."""
    orig_rows = list(csv.DictReader(io.StringIO(orig_text)))
    new_rows = list(csv.DictReader(io.StringIO(new_text)))
    assert len(new_rows) == len(orig_rows)
    newly_assigned = moved = 0
    for old, new in zip(orig_rows, new_rows, strict=True):
        assert {k: v for k, v in new.items() if k != "takt_id"} == {
            k: v for k, v in old.items() if k != "takt_id"
        }
        if old["takt_id"] == new["takt_id"]:
            continue
        assert new["takt_id"], f"element {old['ElementId']} lost its takt zone"
        if old["takt_id"]:
            moved += 1
        else:
            newly_assigned += 1
    assert (newly_assigned, moved) == (TAKT_FIX_NEWLY_ASSIGNED, TAKT_FIX_MOVED)


def _sha256(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


@pytest.mark.parametrize("relative", list(REFERENCE_SHA256))
def test_checkout_has_reference_outputs(ipd_challenge_dir: Path, relative: str) -> None:
    """The checkout is the P0.4 reference state (docs/ROADMAP.md §1)."""
    assert _sha256((ipd_challenge_dir / relative).read_bytes()) == REFERENCE_SHA256[relative]


def test_macro_schedule_reproduces_reference_checksum(runs: dict[str, object]) -> None:
    path = runs["new"] / "alice" / "Macro_Schedule.csv"
    assert _sha256(path.read_bytes()) == REFERENCE_SHA256[f"{ALICE_OUT}/Macro_Schedule.csv"]


def test_takt_schedule_reproduces_reference_checksum(runs: dict[str, object]) -> None:
    path = runs["new"] / "takt" / "Takt_Schedule.csv"
    assert _sha256(path.read_bytes()) == REFERENCE_SHA256[f"{TAKT_OUT}/Takt_Schedule.csv"]


def test_central_bim_model_reproduces_reference_checksum(runs: dict[str, object]) -> None:
    """Identical except for ``source_schedule`` and, since P3B.8 fix 1, ``takt_id``.

    ``source_schedule`` holds each run's absolute input path: the committed file was written
    on the author's machine, so its folder prefix is read from the committed file itself and
    substituted for this run's prefix before hashing. ``takt_id`` changes on purpose (fix 1,
    checked by ``_check_takt_ids_fix``); the committed values are put back before hashing,
    which proves that every other byte is still the reference.
    """
    ipd = runs["ipd"]
    relative = f"{ZONES_OUT}/central_bim_model_with_takt.csv"
    committed = (ipd / relative).read_text(encoding="utf-8")
    match = re.search(r",([^,\n]*[\\/])01_Island_MEP_Concept2_MEP_TakeOff\.csv,", committed)
    assert match, "no source_schedule path found in the committed file"
    current_dir = ipd / "revit_schedules" / "Current"
    regenerated = (runs["new"] / "takt_zones" / "central_bim_model_with_takt.csv").read_text(
        encoding="utf-8"
    )
    regenerated = regenerated.replace(f"{current_dir}{os.sep}", match.group(1))
    regenerated = _with_column_from(regenerated, committed, "takt_id")
    assert _sha256(regenerated.encode("utf-8")) == REFERENCE_SHA256[relative]


def _with_column_from(text: str, source: str, column: str) -> str:
    """``text`` (pandas CSV) with ``column`` replaced by the values of ``source``."""
    rows = list(csv.reader(io.StringIO(text)))
    source_rows = list(csv.reader(io.StringIO(source)))
    index = rows[0].index(column)
    assert source_rows[0].index(column) == index and len(source_rows) == len(rows)
    for row, source_row in zip(rows[1:], source_rows[1:], strict=True):
        row[index] = source_row[index]
    out = io.StringIO()
    csv.writer(out, lineterminator="\n").writerows(rows)
    return out.getvalue()


# P3B.8 fix 2: the committed room_takt_zones.csv was made from an earlier export of the room
# boundaries than the committed one (both in IPD commit c071034). (room_id, column) ->
# (committed, regenerated); every other value is identical.
ROOM_TAKT_ZONES_NEWER_EXPORT = {
    ("1440376", "area_sf"): ("5587.212", "5587.307"),  # L 1 Room 150
    ("1440376", "boundary_segments"): ("223", "241"),
    ("1440382", "boundary_segments"): ("4", "5"),  # L 1 Room 156
    ("1440384", "boundary_segments"): ("4", "6"),  # L 1 Room 158
    ("1440385", "boundary_segments"): ("4", "5"),  # L 1 Room 159
}


def test_room_takt_zones_reproduces_committed_file(runs: dict[str, object]) -> None:
    """The new generator rebuilds IPD's committed room_takt_zones.csv from the export."""
    committed = list(csv.DictReader(
        (runs["ipd"] / ROOMS_OUT / "room_takt_zones.csv").open(encoding="utf-8", newline="")
    ))
    generated = list(csv.DictReader(
        (runs["new"] / "rooms" / "room_takt_zones.csv").open(encoding="utf-8", newline="")
    ))
    assert [row["room_takt_id"] for row in generated] == [
        row["room_takt_id"] for row in committed]
    differences = {
        (old["room_id"], column): (old[column], new[column])
        for old, new in zip(committed, generated, strict=True)
        for column in old
        if old[column] != new[column]
    }
    assert differences == ROOM_TAKT_ZONES_NEWER_EXPORT


@pytest.mark.xfail(
    strict=True,
    reason="The committed Micro_Schedule.csv predates the committed central BIM model "
    "(IPD commit c071034); the original code at 989a6b7 does not reproduce it either. "
    "See engines/schedule/README.md.",
)
def test_micro_schedule_reproduces_reference_checksum(runs: dict[str, object]) -> None:
    path = runs["new"] / "micro" / "Micro_Schedule.csv"
    assert _sha256(path.read_bytes()) == REFERENCE_SHA256[f"{MICRO_OUT}/Micro_Schedule.csv"]
